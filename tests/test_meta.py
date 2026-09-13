import io
import os
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

from PIL import Image
from fastapi.testclient import TestClient

os.environ.setdefault("EXPORT_PASSWORD", "1234")

from app.main import DOWNLOADS_DIR, app
from app.meta import (
    WHERE_FROM_XATTR,
    apply_comment,
    apply_date,
    apply_decrypt,
    apply_hidden,
    apply_scrub,
    apply_tag,
    apply_unlock,
    has_custom_icon,
    is_darwin,
    is_hidden,
    parse_when,
    plant_custom_icon,
    read_comment,
    read_tag,
    read_where_from,
    _gps_present,
    _xattr_list,
    _xattr_set,
    QUARANTINE_XATTR,
)


class ParseWhenTest(unittest.TestCase):
    def test_naive_local(self):
        ts = parse_when("2020-06-15T12:30")
        dt = datetime.fromtimestamp(ts)
        self.assertEqual(dt.year, 2020)
        self.assertEqual(dt.month, 6)
        self.assertEqual(dt.day, 15)
        self.assertEqual(dt.hour, 12)
        self.assertEqual(dt.minute, 30)


class DateOpTest(unittest.TestCase):
    def test_sets_mtime(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "input_photo.png"
            Image.new("RGB", (8, 8), (1, 2, 3)).save(src, "PNG")
            out = apply_date(src, Path(d), modified_at="2020-06-15T12:30")
            want = parse_when("2020-06-15T12:30")
            got = out.stat().st_mtime
            self.assertLess(abs(got - want), 2)
            if is_darwin() and hasattr(out.stat(), "st_birthtime"):
                self.assertLess(abs(out.stat().st_birthtime - want), 2)


class TagOpTest(unittest.TestCase):
    @unittest.skipUnless(is_darwin(), "Finder tags are macOS xattrs")
    def test_roundtrip_red_and_custom(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "input_photo.png"
            Image.new("RGB", (8, 8), (9, 9, 9)).save(src, "PNG")
            out = apply_tag(src, Path(d), "red, work")
            tag = read_tag(out)
            self.assertIn("Red", tag)
            self.assertIn("work", tag)


class UnlockOpTest(unittest.TestCase):
    @unittest.skipUnless(is_darwin(), "quarantine xattr is macOS")
    def test_strips_quarantine(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "input_photo.png"
            Image.new("RGB", (8, 8), (4, 5, 6)).save(src, "PNG")
            from app.meta import _xattr_list, _xattr_set
            _xattr_set(src, QUARANTINE_XATTR, b"0000;test")
            self.assertIn(QUARANTINE_XATTR, _xattr_list(src))
            out = apply_unlock(src, Path(d))
            self.assertNotIn(QUARANTINE_XATTR, _xattr_list(out))


class DecryptOpTest(unittest.TestCase):
    @unittest.skipUnless(shutil.which("zip"), "zip binary missing")
    def test_zip_password(self):
        with tempfile.TemporaryDirectory() as d:
            inner = Path(d) / "secret.txt"
            inner.write_text("hello-export\n", encoding="utf-8")
            zpath = Path(d) / "input_locked.zip"
            r = subprocess.run(
                ["zip", "-P", "s3cret", str(zpath), inner.name],
                cwd=d, capture_output=True, text=True, timeout=15,
            )
            self.assertEqual(r.returncode, 0, r.stderr)
            out = apply_decrypt(zpath, Path(d), "s3cret")
            self.assertIn(b"hello-export", out.read_bytes())

    def test_pdf_password(self):
        import pikepdf
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "input_locked.pdf"
            pdf = pikepdf.new()
            pdf.add_blank_page()
            pdf.save(src, encryption=pikepdf.Encryption(owner="pw", user="pw", R=6))
            out = apply_decrypt(src, Path(d), "pw")
            with pikepdf.open(out) as unlocked:
                self.assertGreaterEqual(len(unlocked.pages), 1)

    def test_7z_password(self):
        import py7zr
        with tempfile.TemporaryDirectory() as d:
            inner = Path(d) / "secret.txt"
            inner.write_text("hello-7z\n", encoding="utf-8")
            zpath = Path(d) / "input_locked.7z"
            with py7zr.SevenZipFile(zpath, "w", password="s3cret") as zf:
                zf.write(inner, arcname="secret.txt")
            out = apply_decrypt(zpath, Path(d), "s3cret")
            self.assertIn(b"hello-7z", out.read_bytes())

    def test_xlsx_password(self):
        import msoffcrypto
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            plain = Path(d) / "plain.xlsx"
            wb = Workbook()
            wb.active["A1"] = "hello-xlsx"
            wb.save(plain)
            enc = Path(d) / "input_locked.xlsx"
            with plain.open("rb") as fh, enc.open("wb") as out:
                msoffcrypto.OfficeFile(fh).encrypt("s3cret", out)
            out = apply_decrypt(enc, Path(d), "s3cret")
            from openpyxl import load_workbook
            self.assertEqual(load_workbook(out).active["A1"].value, "hello-xlsx")

    def test_wrong_password_errors_job(self):
        import py7zr
        client = TestClient(app)
        with tempfile.TemporaryDirectory() as d:
            inner = Path(d) / "secret.txt"
            inner.write_text("x\n", encoding="utf-8")
            zpath = Path(d) / "locked.7z"
            with py7zr.SevenZipFile(zpath, "w", password="right") as zf:
                zf.write(inner, arcname="secret.txt")
            data = zpath.read_bytes()
        r = client.post(
            "/api/meta",
            headers={"X-Password": "1234"},
            files={"file": ("locked.7z", io.BytesIO(data), "application/x-7z-compressed")},
            data={"op": "decrypt", "password": "wrong"},
        )
        self.assertEqual(r.status_code, 200, r.text)
        job_id = r.json()["job_id"]
        st = client.get(f"/api/jobs/{job_id}", headers={"X-Password": "1234"}).json()
        self.assertEqual(st["status"], "error", st)
        self.assertNotEqual(st["status"], "ready")


class MetaApiTest(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_capabilities_meta(self):
        d = self.client.get("/api/capabilities").json()
        self.assertIn("date", d["operations"]["meta_op"])
        self.assertIn("scrub", d["operations"]["meta_op"])
        self.assertIn("comment", d["meta"]["op"])
        self.assertIn("hidden", d["meta"]["op"])
        self.assertIn("tag", d["meta"]["op"])
        self.assertEqual(d["meta"]["is_finder_tag"], is_darwin())
        self.assertEqual(d["meta"]["is_reveal"], is_darwin())

    def test_date_via_api(self):
        buf = io.BytesIO()
        Image.new("RGB", (8, 8), (1, 2, 3)).save(buf, "PNG")
        buf.seek(0)
        r = self.client.post(
            "/api/meta",
            headers={"X-Password": "1234"},
            files={"file": ("photo.png", buf, "image/png")},
            data={"op": "date", "modified_at": "2019-01-02T03:04"},
        )
        self.assertEqual(r.status_code, 200, r.text)
        job_id = r.json()["job_id"]
        st = self.client.get(f"/api/jobs/{job_id}", headers={"X-Password": "1234"}).json()
        self.assertEqual(st["status"], "ready", st)
        dl = self.client.get(f"/api/files/{job_id}", params={"token": st["token"]})
        self.assertEqual(dl.status_code, 200)
        self.assertEqual(dl.content[:8], b"\x89PNG\r\n\x1a\n")

    def test_ui_has_meta_actions(self):
        html = (Path(__file__).resolve().parents[1] / "static" / "index.html").read_text()
        self.assertIn("/api/meta", html)
        self.assertIn("Finder tags", html)
        self.assertIn("show in Finder", html)
        self.assertIn('"date"', html)
        self.assertIn('"scrub"', html)
        self.assertIn('"comment"', html)
        self.assertIn("clear-icon", html)
        self.assertIn("docx", html)
        self.assertIn("acts.push(\"date\", \"tag\")", html)
        self.assertNotIn("?pw=", html)

    def test_unknown_op_400(self):
        buf = io.BytesIO()
        Image.new("RGB", (4, 4), (0, 0, 0)).save(buf, "PNG")
        buf.seek(0)
        r = self.client.post(
            "/api/meta",
            headers={"X-Password": "1234"},
            files={"file": ("a.png", buf, "image/png")},
            data={"op": "explode"},
        )
        self.assertEqual(r.status_code, 400)


class ScrubOpTest(unittest.TestCase):
    def test_jpeg_gps_and_wherefroms_gone_pixels_same(self):
        import piexif
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "input_geo.jpg"
            im = Image.new("RGB", (16, 12), (10, 20, 30))
            gps = {
                piexif.GPSIFD.GPSLatitudeRef: "N",
                piexif.GPSIFD.GPSLatitude: ((37, 1), (48, 1), (0, 1)),
                piexif.GPSIFD.GPSLongitudeRef: "W",
                piexif.GPSIFD.GPSLongitude: ((122, 1), (24, 1), (0, 1)),
            }
            dumped = piexif.dump({
                "0th": {piexif.ImageIFD.Make: "TestCam"},
                "GPS": gps,
                "Exif": {piexif.ExifIFD.BodySerialNumber: b"SN999"},
            })
            im.save(src, "jpeg", exif=dumped)
            before = list(Image.open(src).getdata())
            self.assertTrue(_gps_present(src))
            if is_darwin():
                from app.meta import write_where_from
                write_where_from(src, "https://example.com/photo")
                self.assertTrue(read_where_from(src))
            out = apply_scrub(src, Path(d))
            after = list(Image.open(out).getdata())
            self.assertEqual(before, after)
            self.assertFalse(_gps_present(out))
            self.assertEqual(read_where_from(out), "")
            self.assertNotIn(WHERE_FROM_XATTR, _xattr_list(out))


class CommentOpTest(unittest.TestCase):
    @unittest.skipUnless(is_darwin(), "Finder comment xattrs are macOS")
    def test_set_and_clear(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "input_photo.png"
            Image.new("RGB", (8, 8), (1, 1, 1)).save(src, "PNG")
            out = apply_comment(src, Path(d), "hello", "https://example.com/a")
            self.assertEqual(read_comment(out), "hello")
            self.assertEqual(read_where_from(out), "https://example.com/a")
            src2 = Path(d) / "input_photo2.png"
            shutil.copy2(out, src2)
            cleared = apply_comment(src2, Path(d), "", "")
            self.assertEqual(read_comment(cleared), "")
            self.assertEqual(read_where_from(cleared), "")


class HiddenOpTest(unittest.TestCase):
    @unittest.skipUnless(is_darwin(), "UF_HIDDEN / FinderInfo are macOS")
    def test_hide_then_show(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "input_photo.png"
            Image.new("RGB", (8, 8), (2, 2, 2)).save(src, "PNG")
            hid = apply_hidden(src, Path(d), "hide")
            self.assertTrue(is_hidden(hid))
            src2 = Path(d) / "input_h.png"
            shutil.copy2(hid, src2)
            shown = apply_hidden(src2, Path(d), "show")
            self.assertFalse(is_hidden(shown))

    @unittest.skipUnless(is_darwin(), "FinderInfo custom icon is macOS")
    def test_clear_custom_icon(self):
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "input_photo.png"
            Image.new("RGB", (8, 8), (3, 3, 3)).save(src, "PNG")
            plant_custom_icon(src)
            self.assertTrue(has_custom_icon(src))
            out = apply_hidden(src, Path(d), "clear-icon")
            self.assertFalse(has_custom_icon(out))


class BulkMetaApiTest(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def _png(self, color):
        buf = io.BytesIO()
        Image.new("RGB", (8, 8), color).save(buf, "PNG")
        buf.seek(0)
        return buf

    def test_bulk_date_on_disk_copies(self):
        r = self.client.post(
            "/api/meta",
            headers={"X-Password": "1234"},
            files=[
                ("files", ("a.png", self._png((1, 2, 3)), "image/png")),
                ("files", ("b.png", self._png((4, 5, 6)), "image/png")),
            ],
            data={"op": "date", "modified_at": "2018-03-04T05:06"},
        )
        self.assertEqual(r.status_code, 200, r.text)
        job_id = r.json()["job_id"]
        st = self.client.get(f"/api/jobs/{job_id}", headers={"X-Password": "1234"}).json()
        self.assertEqual(st["status"], "ready", st)
        want = parse_when("2018-03-04T05:06")
        job_dir = DOWNLOADS_DIR / job_id
        copy = [job_dir / "a.png", job_dir / "b.png"]
        for p in copy:
            self.assertTrue(p.exists(), p)
            self.assertLess(abs(p.stat().st_mtime - want), 2)

    @unittest.skipUnless(is_darwin(), "Finder tags are macOS")
    def test_bulk_tag_on_disk_copies(self):
        r = self.client.post(
            "/api/meta",
            headers={"X-Password": "1234"},
            files=[
                ("files", ("a.png", self._png((1, 2, 3)), "image/png")),
                ("files", ("b.png", self._png((4, 5, 6)), "image/png")),
            ],
            data={"op": "tag", "tag": "red, work"},
        )
        self.assertEqual(r.status_code, 200, r.text)
        job_id = r.json()["job_id"]
        st = self.client.get(f"/api/jobs/{job_id}", headers={"X-Password": "1234"}).json()
        self.assertEqual(st["status"], "ready", st)
        job_dir = DOWNLOADS_DIR / job_id
        for name in ("a.png", "b.png"):
            tag = read_tag(job_dir / name)
            self.assertIn("Red", tag)
            self.assertIn("work", tag)

    def test_scrub_via_api(self):
        buf = io.BytesIO()
        Image.new("RGB", (8, 8), (9, 8, 7)).save(buf, "PNG")
        buf.seek(0)
        r = self.client.post(
            "/api/meta",
            headers={"X-Password": "1234"},
            files={"file": ("p.png", buf, "image/png")},
            data={"op": "scrub"},
        )
        self.assertEqual(r.status_code, 200, r.text)
        job_id = r.json()["job_id"]
        st = self.client.get(f"/api/jobs/{job_id}", headers={"X-Password": "1234"}).json()
        self.assertEqual(st["status"], "ready", st)
        dl = self.client.get(f"/api/files/{job_id}", params={"token": st["token"]})
        self.assertEqual(dl.status_code, 200)
        self.assertEqual(dl.content[:8], b"\x89PNG\r\n\x1a\n")
